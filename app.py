import streamlit as st
from openai import OpenAI
import docx
from docx import Document
import io

st.title("🧠 AI Branded Content Integration Agent")

st.markdown("### Step 1: Enter Brand Brief")
brand = st.text_input("Brand")
product = st.text_input("Product")
integration_type = st.radio("Type of Product Integrations Needed", ["1 Active and 1 Passive", "Only Passive", "Only Active"])
communication = st.text_area("Communication (e.g. Tastes like home)")

st.markdown("### Step 2: Upload Show Concept `.docx` Files")
uploaded_files = st.file_uploader("Upload one or more `.docx` files", type="docx", accept_multiple_files=True)

openai_api_key = st.text_input("Enter your OpenAI API Key", type="password")

if st.button("Generate Proposal"):
    if not all([brand, product, integration_type, communication, uploaded_files, openai_api_key]):
        st.error("Please fill out all fields and upload at least one file.")
    else:
        client = OpenAI(api_key=openai_api_key)

        shows = []

        for file in uploaded_files:
            doc = docx.Document(file)
            content = {"Show Name": "", "Themes": "", "Episodes": []}
            for para in doc.paragraphs:
                text = para.text.strip()
                if text.startswith("Show Name-"):
                    content["Show Name"] = text.replace("Show Name-", "").strip()
                elif text.startswith("Themes-"):
                    content["Themes"] = text.replace("Themes-", "").strip()
                elif text.startswith("Episode"):
                    content["Episodes"].append(text)
            shows.append(content)

        proposal_doc = Document()
        proposal_doc.add_heading(f'Branded Integration Proposal for {brand}', 0)

        for show in shows:
            prompt = f"""
You are a creative screenwriter. Given this show:

Show Name: {show['Show Name']}
Themes: {show['Themes']}
Episodes: {', '.join(show['Episodes'][:2])}

And this brand brief:
Brand: {brand}
Product: {product}
Integration Type: {integration_type}
Communication: {communication}

Suggest one cinematic **Active Integration Scene** and one **Passive Integration Scene** for the show. 
Make sure the active scene aligns with the brand communication emotionally. 
Write both scenes in a screenplay-like tone (20
